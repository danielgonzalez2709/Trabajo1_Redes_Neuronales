"""Estadísticas descriptivas de las fuentes del Bloque 1 (T2.1). Sin internet: lee data/raw y data/processed.

Paso 4 del pipeline (orden completo en scripts/comun_bloque1.py).

Uso:
    python scripts/caracterizar_bloque1.py --fecha 2026-10-07
Salida:
    data/processed/caracterizacion_bloque1.json  (y resumen por pantalla)
Las cifras de data/caracterizacion_bloque1.md salen de este archivo.
"""

from __future__ import annotations

import argparse
import collections
import csv
import gzip
import itertools
import json
import statistics

from comun_bloque1 import NGMEP, PROCESSED, RAW, coordenadas_osm, haversine_km, leer_ngmep


def perfil_tabla(filas: list[dict], columnas: list[str]) -> dict:
    """Nº de filas/columnas, nulos (vacío o None) por columna y primer ejemplo no vacío."""
    nulos = {c: sum(1 for f in filas if f.get(c) in ("", None)) for c in columnas}
    ejemplo = {c: next((f[c] for f in filas if f.get(c) not in ("", None)), None) for c in columnas}
    return {"filas": len(filas), "columnas": len(columnas), "nombres_columnas": columnas,
            "nulos": {c: n for c, n in nulos.items() if n}, "ejemplo": ejemplo}


def perfil_ngmep() -> dict:
    out = {}
    for nombre in ["PROVINCIAS.csv", "MUNICIPIOS.csv", "ENTIDADES.csv", "EATIMS.csv", "COMJURIDIC.csv"]:
        crudo = (NGMEP / nombre).read_bytes()
        filas = leer_ngmep(nombre)
        columnas = list(filas[0].keys())
        p = perfil_tabla(filas, columnas)
        p["bytes_0x80_0x9f"] = sorted({b for b in crudo if 0x80 <= b <= 0x9F})
        p["fin_de_linea_crlf"] = b"\r\n" in crudo
        if nombre == "ENTIDADES.csv":
            p["tipos"] = dict(collections.Counter(f["TIPO"] for f in filas))
            p["origencoor"] = dict(collections.Counter(f["ORIGENCOOR"] for f in filas))
            p["codigoine_unico"] = len({f["CODIGOINE"] for f in filas}) == len(filas)
            # ¿La entidad "Capital de municipio" tiene la misma coordenada (texto) que su fila "Municipio"?
            por_codigo = {f["CODIGOINE"]: f for f in filas}
            capitales = [f for f in filas if f["TIPO"] == "Capital de municipio"]
            p["capital_de_municipio_igual_a_municipio"] = sum(
                (f["LATITUD_ETRS89_REGCAN95"], f["LONGITUD_ETRS89_REGCAN95"])
                == (por_codigo[f["INEMUNI"] + "000000"]["LATITUD_ETRS89_REGCAN95"],
                    por_codigo[f["INEMUNI"] + "000000"]["LONGITUD_ETRS89_REGCAN95"])
                for f in capitales)
            p["capitales_de_municipio"] = len(capitales)
            p["suprimida_ine_verdadero"] = sum(f["SUPRIMIDA_INE"] == "VERDADERO" for f in filas)
            p["discrepante_ine_verdadero"] = sum(f["DISCREPANTE_INE"] == "VERDADERO" for f in filas)
        out[nombre] = p
    return out


def perfil_nuc(fecha: str) -> dict:
    consultas = json.loads((RAW / f"ign_nuc_capitales_provincia_{fecha}.json").read_text(encoding="utf-8"))
    feats = [f["properties"] for c in consultas for f in c["respuesta"]["features"]]
    columnas = sorted({k for f in feats for k in f})
    p = perfil_tabla(feats, columnas)
    p["consultas"] = len(consultas)
    p["consultas_con_resultado"] = sum(1 for c in consultas if c["respuesta"]["features"])
    p["codigos_capital"] = dict(collections.Counter(f["capital"] for f in feats))
    p["tipos"] = dict(collections.Counter(f["tipo"] for f in feats))
    p["fecha_registro_min"] = min(f["fecha"][:10] for f in feats)
    p["fecha_registro_max"] = max(f["fecha"][:10] for f in feats)
    return p


def perfil_provincias(fecha: str) -> dict:
    with gzip.open(RAW / f"ign_provincias_{fecha}.geojson.gz", "rt", encoding="utf-8") as fh:
        datos = json.load(fh)
    props = [f["properties"] for f in datos["features"]]
    columnas = sorted({k for f in props for k in f})
    p = perfil_tabla(props, columnas)
    p["tipos_geometria"] = dict(collections.Counter(f["geometry"]["type"] for f in datos["features"]))
    p["vertices_total"] = sum(len(anillo) for f in datos["features"]
                              for pol in (f["geometry"]["coordinates"] if f["geometry"]["type"] == "MultiPolygon"
                                          else [f["geometry"]["coordinates"]])
                              for anillo in pol)
    p["timeStamp"] = datos.get("timeStamp")
    p["bytes_gz"] = (RAW / f"ign_provincias_{fecha}.geojson.gz").stat().st_size
    return p


def perfil_ine(fecha: str) -> dict:
    import openpyxl

    libro = openpyxl.load_workbook(RAW / f"ine_26codmun_{fecha}.xlsx", read_only=True)
    municipios = []
    for hoja in libro.worksheets:
        filas = list(hoja.iter_rows(values_only=True))
        cab = list(filas[2])
        municipios += [dict(zip(cab, f)) for f in filas[3:] if f[0]]
    p = perfil_tabla(municipios, ["CPRO", "CMUN", "DC", "NOMBRE"])
    p["hojas"] = len(libro.worksheets)
    p["titulo"] = list(libro.worksheets[0].iter_rows(values_only=True))[0][0]
    return p


def perfil_nominatim(fecha: str) -> dict:
    consultas = json.loads((RAW / f"nominatim_capitales_{fecha}.json").read_text(encoding="utf-8"))
    primeros = [c["respuesta"][0] for c in consultas if c["respuesta"] and c["modo"] == "es"]
    return {
        "consultas": len(consultas),
        "por_modo": dict(collections.Counter(c["modo"] for c in consultas)),
        "resultados_por_consulta": dict(collections.Counter(len(c["respuesta"]) for c in consultas)),
        "addresstype_primer_resultado_es": dict(collections.Counter(r.get("addresstype") for r in primeros)),
        "pais_primer_resultado_libre": {c["consulta"]: (c["respuesta"][0].get("address") or {}).get("country_code")
                                        for c in consultas if c["modo"] == "libre"},
        "columnas_resultado": sorted(primeros[0].keys()),
        "ejemplo": {k: primeros[0][k] for k in ("place_id", "osm_type", "lat", "lon", "addresstype", "display_name")},
    }


def perfil_cities() -> dict:
    with (PROCESSED / "cities.csv").open(encoding="utf-8", newline="") as fh:
        lector = csv.DictReader(fh)
        filas = list(lector)
        columnas = list(lector.fieldnames or [])
    p = perfil_tabla(filas, columnas)
    lat = [float(f["lat"]) for f in filas]
    lon = [float(f["lon"]) for f in filas]
    p["lat_min_max"] = [min(lat), max(lat)]
    p["lon_min_max"] = [min(lon), max(lon)]
    p["verificado_manual"] = dict(collections.Counter(f["verificado_manual"] for f in filas))
    pares = sorted((haversine_km(float(a["lat"]), float(a["lon"]), float(b["lat"]), float(b["lon"])), a["capital"], b["capital"])
                   for a, b in itertools.combinations(filas, 2))
    p["par_mas_cercano_km"] = [round(pares[0][0], 3), pares[0][1], pares[0][2]]
    p["par_mas_lejano_km"] = [round(pares[-1][0], 3), pares[-1][1], pares[-1][2]]
    with (PROCESSED / "cities_verificacion.csv").open(encoding="utf-8", newline="") as fh:
        v = list(csv.DictReader(fh))
    for r in v:
        coordenadas_osm(r)  # mismo criterio que construir/validar: la falta de OSM es un error explícito
    d_osm = [float(r["dist_osm_km"]) for r in v]
    p["osm_dist_km"] = {"max": max(d_osm), "mediana": statistics.median(d_osm), "mayores_3km":
                        sorted([(r["capital"], float(r["dist_osm_km"])) for r in v if float(r["dist_osm_km"]) > 3],
                               key=lambda t: -t[1])}
    p["en_provincia_ign"] = dict(collections.Counter(r["en_provincia_ign"] for r in v))
    p["fuente"] = dict(collections.Counter(f["fuente"] for f in filas))
    p["fecha_consulta"] = dict(collections.Counter(f["fecha_consulta"] for f in filas))
    p["origencoor_ngmep"] = dict(collections.Counter(r["origencoor_ngmep"] for r in v))
    entidades = {e["CODIGOINE"]: e for e in leer_ngmep("ENTIDADES.csv")}
    coord = lambda e: (e["LATITUD_ETRS89_REGCAN95"], e["LONGITUD_ETRS89_REGCAN95"])  # noqa: E731
    p["capitales_con_coordenada_igual_a_su_municipio_ngmep"] = sum(
        coord(entidades[r["codigoine"]]) == coord(entidades[r["codigoine"][:5] + "000000"]) for r in v)
    p["origencoor_ngmep_de_las_distintas_a_igr"] = dict(collections.Counter(
        r["origencoor_ngmep"] for r in v if float(r["dist_igr_km"]) >= 0.001))
    contraste = PROCESSED / "cities_contraste_ngmep.csv"
    if contraste.exists():
        with contraste.open(encoding="utf-8", newline="") as fh:
            c = list(csv.DictReader(fh))
        d = [float(r["dist_km"]) for r in c]
        p["ngmep"] = {"max_km": max(d), "mediana_km": statistics.median(d), "media_km": round(statistics.mean(d), 3),
                      "identicas": sum(x < 0.001 for x in d),
                      "codigos_coinciden": sum(r["codine_coincide"] == "true" for r in c),
                      "mayores_0_5km": [(r["capital"], r["cod_prov"], float(r["dist_km"])) for r in c
                                        if float(r["dist_km"]) > 0.5],
                      "distintas_por_cod_prov": [r["cod_prov"] for r in c if float(r["dist_km"]) >= 0.001],
                      "distintas_ngmep_mas_cerca_de_osm": [r["capital"] for r in c if float(r["dist_km"]) >= 0.001
                                                           and float(r["dist_osm_ngmep_km"]) < float(r["dist_osm_ign_km"])],
                      "distintas_igr_mas_cerca_de_osm": [r["capital"] for r in c if float(r["dist_km"]) >= 0.001
                                                         and float(r["dist_osm_ngmep_km"]) > float(r["dist_osm_ign_km"])]}
    return p


def construir_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--fecha", required=True,
                        help="Fecha de descarga de los crudos automáticos en data/raw/ (AAAA-MM-DD)")
    return parser


def main() -> None:
    args = construir_parser().parse_args()
    resultado = {
        "ngmep": perfil_ngmep(),
        "ign_nuc": perfil_nuc(args.fecha),
        "ign_provincias": perfil_provincias(args.fecha),
        "ine_codmun": perfil_ine(args.fecha),
        "nominatim": perfil_nominatim(args.fecha),
        "cities": perfil_cities(),
    }
    destino = PROCESSED / "caracterizacion_bloque1.json"
    destino.write_text(json.dumps(resultado, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
    print(json.dumps(resultado, ensure_ascii=False, indent=1, default=str))


if __name__ == "__main__":
    main()
